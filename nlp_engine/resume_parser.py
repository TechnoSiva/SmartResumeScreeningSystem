"""
Resume Parser — extracts raw text and contact info from PDF / DOCX files.
"""
import re
from pathlib import Path

from typing import Optional

import pdfplumber
from docx import Document
import spacy

from config import SPACY_MODEL

# Load spaCy model lazily
_nlp = None

def _get_nlp():
    global _nlp
    if _nlp is None:
        _nlp = spacy.load(SPACY_MODEL)
    return _nlp


def extract_text_from_pdf(file_path: str) -> str:
    """Extract text from a PDF file."""
    text_parts = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
    return "\n".join(text_parts)


def extract_text_from_docx(file_path: str) -> str:
    """Extract text from a DOCX file."""
    doc = Document(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    # Also extract from tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text.strip())
    return "\n".join(paragraphs)


def extract_text(file_path: str) -> str:
    """Extract text from a resume file (PDF or DOCX)."""
    path = Path(file_path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext == ".docx":
        return extract_text_from_docx(file_path)
    else:
        raise ValueError(f"Unsupported file format: {ext}")


def extract_email(text: str) -> Optional[str]:
    """Extract email address from text."""
    pattern = r"[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}"
    match = re.search(pattern, text)
    return match.group(0) if match else None


def extract_phone(text: str) -> Optional[str]:
    """Extract phone number from text."""
    patterns = [
        r"(?:\+?\d{1,3}[\s\-]?)?\(?\d{3}\)?[\s\-]?\d{3}[\s\-]?\d{4}",
        r"(?:\+?\d{1,3}[\s\-]?)?\d{5}[\s\-]?\d{5}",
        r"(?:\+?\d{1,2}[\s\-])?\d{10}",
    ]
    for pattern in patterns:
        match = re.search(pattern, text)
        if match:
            return match.group(0).strip()
    return None


def extract_name_from_filename(file_path: str) -> Optional[str]:
    """Extract candidate name from the file name if it follows standard conventions."""
    if not file_path:
        return None
        
    stem = Path(file_path).stem.lower()
    
    stop_words = [
        "resume", "cv", "curriculum", "vitae", "profile", "summary", "biodata",
        "updated", "final", "latest", "new", "doc", "docx", "pdf", 
        "fresher", "experienced", "engineer", "developer", "java", "python",
        "intern", "frontend", "backend", "fullstack", "data", "scientist",
        "application", "candidate", "v1", "v2", "v3", "copy"
    ]
    
    # Explicitly remove stop words even if concatenated (e.g., "AcharyResume" -> "Achary")
    for word in stop_words:
        stem = stem.replace(word, " ")
        
    # Replace anything that isn't a letter with a space
    clean = re.sub(r'[^a-zÀ-ÿ]', ' ', stem)
    words = clean.split()
    
    clean_words = []
    for w in words:
        if len(w) > 0:
            clean_words.append(w.capitalize())
            
    # A reliable filename name usually has 2-4 words
    if 2 <= len(clean_words) <= 4:
        return " ".join(clean_words)
    return None


def extract_name(text: str, file_path: str = "") -> str:
    """
    Extract candidate name using filename first, then text heuristics 
    optimized for real-world resumes, with a fallback to spaCy NER.
    """
    # 1. Try filename (very accurate for "First_Last_Resume.pdf")
    name_from_file = extract_name_from_filename(file_path)
    if name_from_file:
        return name_from_file

    lines = [line.strip() for line in text.split("\n") if line.strip()]
    if not lines:
        return "Unknown"
        
    invalid_keywords = {
        "resume", "curriculum", "vitae", "cv", "profile", "summary", "objective", 
        "address", "phone", "email", "mobile", "ph", "e-mail", "dob", "date", 
        "linkedin", "github", "portfolio", "page", "skills", "experience", "education"
    }
    
    for line in lines[:20]:
        line_clean = line
        
        # Strip "Name:" prefix if present
        if line_clean.lower().startswith("name:"):
            line_clean = line_clean[5:].strip()
            
        line_lower = line_clean.lower()
        
        # Skip if it contains email, url, phone numbers, or typical contact indicators
        if re.search(r"@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+|http|www|\.com|\.in|\.org|\d{4,}|\.net|\.io", line_lower):
            continue
            
        # Skip if it contains common non-name resume section headers or words
        tokens = set(re.findall(r'\w+', line_lower))
        if tokens.intersection(invalid_keywords) or "curriculum vitae" in line_lower:
            continue
            
        # A name usually has 2-5 words and is fairly short
        # We require at least 2 words to avoid picking up single skills like "Java"
        words = line_clean.split()
        if 2 <= len(words) <= 5 and len(line_clean) < 40:
            # Check if it consists strictly of characters typical in names
            if re.match(r"^[A-Za-zÀ-ÿ\s\.\-\,\(\)]+$", line_clean):
                if line_clean.isupper():
                    return line_clean.title()
                return line_clean

    # Fallback to spaCy on the first few lines if the heuristic fails
    try:
        nlp = _get_nlp()
        first_lines = "\n".join(lines[:8])
        doc = nlp(first_lines)
        for ent in doc.ents:
            if ent.label_ == "PERSON":
                return ent.text.strip()
    except Exception:
        pass
        
    # Final fallback: return the first line that isn't glaringly contact info
    for line in lines[:5]:
        if len(line) < 60 and not re.search(r"@|http|www|\d{5,}", line):
            return line
            
    return "Unknown"

def parse_resume(file_path: str) -> dict:
    """
    Full resume parsing pipeline.
    Returns dict with: name, email, phone, raw_text
    """
    raw_text = extract_text(file_path)
    return {
        "name": extract_name(raw_text, file_path),
        "email": extract_email(raw_text),
        "phone": extract_phone(raw_text),
        "raw_text": raw_text,
    }
